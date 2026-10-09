"""Streamlit QC report aggregator: upload HTML, preview, and export Excel."""
from __future__ import annotations

import ast
from io import BytesIO
import math
from pathlib import PurePosixPath
import re

import pandas as pd

REPORT_PATTERN = re.compile(
    r"^(?P<flowcell>V\d+)_(?P<lane>L\d{2})"
    r"(?:(?:\.summaryReport)|(?:_(?P<barcode>\d+)\.report))\.html$",
    re.IGNORECASE,
)
METRICS = {"%Q30": "Q30(%)", "%Q40": "Q40(%)", "Reads (M)": "TotalReads(M)"}
DEFAULT_LANES = ["L01", "L02", "L03", "L04"]


def read_used_barcodes(content: bytes) -> set[int]:
    """Read a CSV whitelist with a Barcode column, one used ID per row."""
    try:
        table = pd.read_csv(BytesIO(content), encoding="utf-8-sig", dtype=str)
    except (pd.errors.ParserError, pd.errors.EmptyDataError, UnicodeError) as exc:
        raise ValueError(f"Could not read used-barcode CSV: {exc}") from exc
    columns = [column for column in table if column.strip().casefold() == "barcode"]
    if len(columns) != 1:
        raise ValueError("Used-barcode CSV needs exactly one 'Barcode' column.")
    values = table[columns[0]].dropna().astype(str).str.strip()
    if values.empty or not values.map(lambda value: bool(re.fullmatch(r"\d+", value))).all():
        raise ValueError("Barcode column must contain nonempty integer IDs, such as 1, 5, 12.")
    return {int(value) for value in values}


def parse_report(name: str, content: bytes) -> dict:
    """Extract literal report data only; never execute uploaded JavaScript."""
    path = PurePosixPath(name.replace("\\", "/"))
    match = REPORT_PATTERN.fullmatch(path.name)
    if not match:
        raise ValueError("Filename does not match a lane or barcode report pattern.")
    flowcell, lane = match["flowcell"].upper(), match["lane"].upper()
    barcode = int(match["barcode"]) if match["barcode"] is not None else None
    # Directory metadata is optional: some browser uploads retain only basenames.
    for part in path.parts[:-1]:
        if re.fullmatch(r"V\d+", part, re.IGNORECASE) and part.upper() != flowcell:
            raise ValueError("Flowcell folder disagrees with the filename.")
        if re.fullmatch(r"L\d{2}", part, re.IGNORECASE) and part.upper() != lane:
            raise ValueError("Lane folder disagrees with the filename.")
    try:
        html = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("Report must be UTF-8 encoded.") from exc
    title = re.search(r'''\bvar\s+reportTitle\s*=\s*['"]Analysis Report of (V\d+_L\d{2}(?:_\d+)?)['"]''', html, re.IGNORECASE)
    expected_id = f"{flowcell}_{lane}" + (f"_{barcode}" if barcode is not None else "")
    if title:
        title_match = re.fullmatch(r"(V\d+)_(L\d{2})(?:_(\d+))?", title[1], re.IGNORECASE)
        internal_id = f"{title_match[1].upper()}_{title_match[2].upper()}"
        if title_match[3] is not None:
            internal_id += f"_{int(title_match[3])}"
        if internal_id != expected_id:
            raise ValueError("Report title disagrees with the filename.")
    table_match = re.search(r"\bvar\s+summaryTable\s*=\s*(\[.*?\])\s*;", html, re.DOTALL)
    if not table_match:
        raise ValueError("summaryTable was not found.")
    if len(table_match[1]) > 65536:
        raise ValueError("summaryTable is unexpectedly large.")
    try:
        rows = ast.literal_eval(table_match[1])
    except (ValueError, SyntaxError, RecursionError) as exc:
        raise ValueError("summaryTable is not a supported literal array.") from exc
    if not isinstance(rows, list) or not rows:
        raise ValueError("summaryTable must contain a list of rows.")
    metrics = {}
    for row in rows:
        if not isinstance(row, (list, tuple)) or len(row) != 2 or not isinstance(row[0], str):
            raise ValueError("summaryTable contains an invalid row.")
        if row[0] in metrics:
            raise ValueError(f"Repeated summary metric: {row[0]}")
        metrics[row[0]] = row[1]
    record = {
        "Flowcell": flowcell, "Lane": lane, "Barcode": barcode,
        "Report Type": "Lane" if barcode is None else "Barcode", "Source": name,
    }
    required = dict(METRICS)
    if barcode is None:
        required["Chip Productivity (%)"] = "ChipProductivity(%)"
    for column, key in required.items():
        if key not in metrics:
            raise ValueError(f"Required metric is missing: {key}")
        try:
            value = float(metrics[key])
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Metric is not numeric: {key}") from exc
        if not math.isfinite(value) or value < 0:
            raise ValueError(f"Metric must be finite and nonnegative: {key}")
        if key != "TotalReads(M)" and value > 100:
            raise ValueError(f"Percentage is outside 0–100: {key}")
        record[column] = value
    if record["%Q40"] > record["%Q30"]:
        raise ValueError("Q40 cannot exceed Q30.")
    return record


def collect_reports(files) -> tuple[list[dict], list[dict]]:
    """Parse independently; exclude every member of an ambiguous duplicate group."""
    records, issues = [], []
    for name, content in files:
        try:
            records.append(parse_report(name, content))
        except ValueError as exc:
            issues.append({"Source": name, "Issue": str(exc)})
    groups = {}
    for record in records:
        key = (record["Flowcell"], record["Lane"], record["Barcode"])
        groups.setdefault(key, []).append(record)
    accepted = []
    for group in groups.values():
        if len(group) == 1:
            accepted.extend(group)
        else:
            for record in group:
                issues.append({"Source": record["Source"], "Issue": "Duplicate report identity; all copies excluded. Remove duplicates and retry."})
    return accepted, issues


def build_tables(records: list[dict], flowcell: str, lanes: list[str], barcodes: list[int]):
    selected = [r for r in records if r["Flowcell"] == flowcell and r["Lane"] in lanes]
    lane_records = {r["Lane"]: r for r in selected if r["Report Type"] == "Lane"}
    lane_rows = []
    for lane in lanes:
        record = lane_records.get(lane, {})
        lane_rows.append({"Lane": lane, "%Q30": record.get("%Q30"), "%Q40": record.get("%Q40"),
                          "Total Reads (M)": record.get("Reads (M)"),
                          "Chip Productivity (%)": record.get("Chip Productivity (%)")})
    lane_df = pd.DataFrame(lane_rows).set_index("Lane")
    sample_records = {(r["Barcode"], r["Lane"]): r for r in selected if r["Report Type"] == "Barcode"}
    columns = [(lane, metric) for lane in lanes for metric in METRICS]
    columns += [("Total", "Reads (M)"), ("Reports", "Coverage")]
    barcode_rows, coverage_rows = [], []
    for barcode in sorted(barcodes):
        values, reads, missing = [], [], []
        for lane in lanes:
            record = sample_records.get((barcode, lane))
            values.extend(record[metric] if record else None for metric in METRICS)
            if record:
                reads.append(record["Reads (M)"])
            else:
                missing.append(lane)
        # Available totals are explicitly labelled as partial by Coverage.
        values += [round(math.fsum(reads), 2) if reads else None, f"{len(reads)}/{len(lanes)}"]
        barcode_rows.append(values)
        coverage_rows.append({"Barcode": barcode, "Coverage": f"{len(reads)}/{len(lanes)}",
                              "Status": "Complete" if not missing else "Incomplete",
                              "Missing Lanes": ", ".join(missing)})
    barcode_df = pd.DataFrame(barcode_rows, index=sorted(barcodes), columns=pd.MultiIndex.from_tuples(columns))
    barcode_df.index.name = "Barcode"
    coverage_df = pd.DataFrame(coverage_rows, columns=["Barcode", "Coverage", "Status", "Missing Lanes"])
    return lane_df, barcode_df, coverage_df


def flat_barcodes(df):
    result = df.copy()
    result.columns = [f"{lane} {metric}" for lane, metric in result.columns]
    return result


def excel_bytes(lane_df, barcode_df) -> bytes:
    """Put the two QC tables into separate tabs of one Excel workbook."""
    from openpyxl.styles import Alignment, Font, PatternFill
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        lane_df.to_excel(writer, sheet_name="Lane QC")
        barcode_df.to_excel(writer, sheet_name="Barcode QC", merge_cells=True)
        for ws in writer.book.worksheets:
            header_rows = 2 if ws.title == "Barcode QC" else 1
            ws.freeze_panes = "B4" if ws.title == "Barcode QC" else "B2"
            for row in ws.iter_rows():
                for cell in row:
                    if cell.row <= header_rows:
                        cell.fill = PatternFill("solid", fgColor="173F5F")
                        cell.font = Font(color="FFFFFF", bold=True)
                        cell.alignment = Alignment(horizontal="center")
                    elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        cell.number_format = "0" if cell.column == 1 and ws.title == "Barcode QC" else "0.00"
                    # Preserve untrusted source names as literal text, never formulas.
                    if cell.data_type == "f":
                        cell.data_type = "s"
            for col in ws.columns:
                from openpyxl.utils import get_column_letter
                width = min(65, max(14, max(len(str(c.value or "")) for c in col) + 2))
                ws.column_dimensions[get_column_letter(col[0].column)].width = width
    return output.getvalue()


def run_app():
    import streamlit as st
    st.set_page_config(page_title="Flowcell QC", page_icon="🧬", layout="wide")
    st.title("🧬 Flowcell QC")
    st.write("Upload lane and barcode HTML reports to build your QC tables.")
    mode = st.radio(
        "Report source",
        ["Upload flowcell folder", "Upload HTML files"],
        horizontal=True,
    )
    uploads = st.file_uploader(
        "Upload QC HTML reports", type=["html"],
        accept_multiple_files="directory" if mode == "Upload flowcell folder" else True,
        help="When selecting a folder, only .html files are uploaded; FASTQ files stay on your device.",
        key=f"reports_{mode}",
    )
    if not uploads:
        st.info("Choose reports to preview the QC tables.")
        return
    payload = tuple((file.name, file.getvalue()) for file in uploads)
    @st.cache_data(show_spinner=False)
    def cached_parse(payload):
        return collect_reports(payload)
    records, issues = cached_parse(payload)
    if issues:
        st.warning(f"{len(issues)} file(s) excluded. Review the issues below.")
        with st.expander("Upload issues", expanded=True):
            st.dataframe(pd.DataFrame(issues), hide_index=True, use_container_width=True)
    if not records:
        st.error("No valid reports remain. Correct the issues and upload again.")
        return
    flowcells = sorted({r["Flowcell"] for r in records})
    flowcell = st.selectbox("Flowcell ID", flowcells)
    current = [r for r in records if r["Flowcell"] == flowcell]
    discovered_lanes = sorted({r["Lane"] for r in current})
    lane_options = sorted(set(DEFAULT_LANES + discovered_lanes))
    lanes = st.multiselect("Expected lanes", lane_options, default=lane_options, key=f"lanes_{flowcell}")
    if not lanes:
        st.info("Select at least one expected lane.")
        return
    lanes = sorted(lanes)
    available_barcodes = sorted({r["Barcode"] for r in current if r["Barcode"] is not None and r["Lane"] in lanes})
    choose = st.radio(
        "Which barcodes were used?",
        ["All detected barcodes (no used filter)", "Choose barcodes", "Upload used-barcode CSV"],
        horizontal=True,
    )
    if choose == "All detected barcodes (no used filter)":
        barcodes = available_barcodes
    elif choose == "Choose barcodes":
        barcodes = st.multiselect(
            "Barcodes to include", available_barcodes, default=available_barcodes,
            key=f"barcodes_{flowcell}_{','.join(lanes)}",
        )
    else:
        barcode_file = st.file_uploader(
            "Used-barcode CSV (column: Barcode)", type=["csv"], key=f"used_barcodes_{flowcell}"
        )
        if barcode_file is None:
            st.info("Upload a small CSV containing only the used barcode IDs. Example: Barcode followed by 1, 5, 12 on separate rows.")
            return
        try:
            used = read_used_barcodes(barcode_file.getvalue())
        except ValueError as exc:
            st.error(str(exc))
            return
        absent = sorted(used - set(available_barcodes))
        if absent:
            st.warning(f"Used barcodes without a matching report in selected lanes: {absent}")
        barcodes = sorted(used & set(available_barcodes))
        st.caption(f"{len(barcodes)} used barcode(s) matched; {len(available_barcodes) - len(barcodes)} detected barcode(s) excluded.")
    lane_df, barcode_df, coverage = build_tables(records, flowcell, lanes, barcodes)
    c1, c2, c3 = st.columns(3)
    c1.metric("Valid reports in flowcell", len(current))
    c2.metric("Detected lanes", len(discovered_lanes))
    c3.metric("Selected barcodes", len(barcodes))
    missing_summaries = [lane for lane in lanes if not any(r["Report Type"] == "Lane" and r["Lane"] == lane for r in current)]
    if missing_summaries:
        st.warning("Missing lane summaries: " + ", ".join(missing_summaries))
    st.subheader("Lane QC")
    st.dataframe(lane_df.style.format("{:.2f}", na_rep="Missing"), use_container_width=True)
    st.subheader("Barcode QC")
    if not barcodes:
        st.info("No barcodes selected or no barcode reports available in the selected lanes.")
    else:
        formats = {col: "{:.2f}" for col in barcode_df.columns if col != ("Reports", "Coverage")}
        st.dataframe(barcode_df.style.format(formats, na_rep="Missing"), use_container_width=True)
        incomplete = coverage[coverage["Status"] == "Incomplete"]
        if not incomplete.empty:
            st.warning("Some barcode reports are missing. Total reads sum available reports only; these totals are partial.")
            st.dataframe(incomplete, hide_index=True, use_container_width=True)
    st.caption("Reads are in millions, exactly as reported in summaryTable. No QC thresholds have been applied.")
    st.subheader("Download")
    d1, d2, d3 = st.columns(3)
    d1.download_button("Lane CSV", lane_df.to_csv().encode("utf-8-sig"), f"{flowcell}_lane_qc.csv", "text/csv")
    d2.download_button("Barcode CSV", flat_barcodes(barcode_df).to_csv().encode("utf-8-sig"), f"{flowcell}_barcode_qc.csv", "text/csv")
    d3.download_button("Excel workbook", excel_bytes(lane_df, barcode_df),
                       f"{flowcell}_QC.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    run_app()